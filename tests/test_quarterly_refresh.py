import json
import re
import sys
import unittest
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import build_dashboard as build


class QuarterlyRefreshTests(unittest.TestCase):
    def test_new_period_replaces_overlap_including_missing_old_categories(self):
        def payload(rows):
            return {'records': rows, 'metadata': {'source_file': 'test.xlsx', 'sheet_info': []}}
        old = payload([{'sheet': 'A', 'date': day, 'total_count': 100}
                       for day in ['2026-06-30', '2026-07-01', '2026-07-02']])
        new = payload([{'sheet': 'A', 'date': '2026-07-01', 'total_count': 2}])
        result = build.supplement_history(new, old)
        self.assertEqual([r['total_count'] for r in result['records']], [100, 2])

    def test_export_matches_daily_excel_grand_totals(self):
        html = (ROOT / 'output/packaging_dashboard.html').read_text(encoding='utf-8')
        payload = json.loads(re.search(r'<script[^>]*id="dashboard-data"[^>]*>(.*?)</script>', html, re.S).group(1))
        actual = defaultdict(float)
        for row in payload['records']:
            actual[row['sheet'], row['date']] += row['total_count']
        expected = {}
        # New workbook owns the complete period from each sheet's first date.
        for name in dict.fromkeys(['Data_pro _balení dashboard_6_2026.xlsx',
                                   payload['metadata']['history_source_file'],
                                   payload['metadata']['source_file']]):
            workbook = load_workbook(ROOT / 'input' / name, data_only=True)
            for sheet in workbook:
                header = build.find_pivot_header_row(sheet)
                dates = [(c, build.as_date(sheet.cell(header, c).value))
                         for c in range(2, sheet.max_column + 1)]
                dates = [(c, d) for c, d in dates if re.fullmatch(r'\d{4}-\d{2}-\d{2}', d)]
                start = min(d for _, d in dates)
                expected = {k: v for k, v in expected.items() if k[0] != sheet.title or k[1] < start}
                total_row = next(r for r in range(header + 1, sheet.max_row + 1)
                                 if build.normalize_text(sheet.cell(r, 1).value) in {'celkovy soucet', 'grand total'})
                for col, day in dates:
                    value = build.safe_number(sheet.cell(total_row, col).value)
                    if value:
                        expected[sheet.title, day] = value
            workbook.close()
        self.assertEqual(dict(actual), expected)
        for sheet in ['SKLC3', 'CZLC4']:
            self.assertEqual(max(d for s, d in actual if s == sheet),
                             max(d for s, d in expected if s == sheet))


if __name__ == '__main__':
    unittest.main()
